"""内閣府DIコネクタ（消費動向調査・景気ウォッチャー調査）。

公表 Excel を openpyxl で読み、各DI系列ごとに ``date``(月初)/``value`` の
DataFrame と出典メタ ``Source`` を返す。

代表フィクスチャ（``data/fixtures/cao_*.xlsx``）の列ラベルを基準にした
``column_map``（列ラベル -> series_id）で抽出する。実 Excel のレイアウトが
確定したら column_map / sheet 指定を合わせて差し替える（M1-5 注記）。
季節調整値を採用（フィクスチャ・実データとも季調値列を対象とする）。
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from etl.provenance import Source
from etl.registry import REGISTRY

# 列ラベル -> 内部 series_id
CCI_COLUMN_MAP: dict[str, str] = {
    "消費者態度指数": "cao.cci.attitude",
    "暮らし向き": "cao.cci.livelihood",
    "収入の増え方": "cao.cci.income",
    "雇用環境": "cao.cci.employment",
    "耐久消費財の買い時判断": "cao.cci.durables",
}

WATCHER_COLUMN_MAP: dict[str, str] = {
    "現状判断DI": "cao.watcher.current",
    "先行き判断DI": "cao.watcher.outlook",
}

_DATE_HEADERS = {"年月", "月", "調査年月", "時期"}
_NAME_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月")


def _to_date(cell: Any) -> pd.Timestamp | None:
    if cell is None:
        return None
    if isinstance(cell, datetime):
        return pd.Timestamp(year=cell.year, month=cell.month, day=1)
    m = _NAME_RE.search(str(cell))
    if m:
        return pd.Timestamp(year=int(m.group(1)), month=int(m.group(2)), day=1)
    return None


def _to_float(cell: Any) -> float:
    if cell is None:
        return float("nan")
    if isinstance(cell, (int, float)):
        return float(cell)
    text = str(cell).strip().replace(",", "")
    if text in {"", "-", "***", "…", "...", "X", "x"}:
        return float("nan")
    try:
        return float(text)
    except ValueError:
        return float("nan")


def parse_cao_excel(
    content: bytes,
    column_map: dict[str, str],
    *,
    sheet: str | None = None,
) -> dict[str, pd.DataFrame]:
    """内閣府 Excel を {series_id: DataFrame[date,value]} に変換する。

    ヘッダ行は「``column_map`` のラベルを少なくとも1つ含む行」として検出する。
    日付列はヘッダ行内の日付見出し（年月 等）、無ければ最左列を採用する。
    """
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active

    grid = [list(row) for row in ws.iter_rows(values_only=True)]

    header_idx = None
    for i, row in enumerate(grid):
        labels = {str(c).strip() for c in row if c is not None}
        if labels & set(column_map.keys()):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(
            f"header row not found for columns {list(column_map.keys())}"
        )

    header = [str(c).strip() if c is not None else "" for c in grid[header_idx]]

    # 日付列インデックスの決定。
    date_col = next((j for j, h in enumerate(header) if h in _DATE_HEADERS), 0)

    # ラベル列インデックス。
    label_cols = {
        column_map[h]: j for j, h in enumerate(header) if h in column_map
    }

    result: dict[str, list[tuple[pd.Timestamp, float]]] = {
        sid: [] for sid in label_cols
    }
    for row in grid[header_idx + 1 :]:
        if date_col >= len(row):
            continue
        date = _to_date(row[date_col])
        if date is None:
            continue
        for sid, j in label_cols.items():
            value = _to_float(row[j]) if j < len(row) else float("nan")
            result[sid].append((date, value))

    frames: dict[str, pd.DataFrame] = {}
    for sid, rows in result.items():
        df = pd.DataFrame(rows, columns=["date", "value"])
        df = df.sort_values("date").drop_duplicates("date").reset_index(drop=True)
        df["date"] = pd.to_datetime(df["date"])
        df["value"] = df["value"].astype(float)
        frames[sid] = df
    return frames


def _source_for(series_id: str, retrieved_at: datetime) -> Source:
    spec = REGISTRY[series_id]
    return Source(
        series_id=spec.series_id,
        name=spec.name,
        url=spec.url,
        retrieved_at=retrieved_at,
        license=spec.license,
        unit=spec.unit,
        frequency=spec.frequency,
    )


def _fetch(
    content: bytes,
    column_map: dict[str, str],
    *,
    sheet: str | None,
    retrieved_at: datetime | None,
) -> list[tuple[pd.DataFrame, Source]]:
    when = retrieved_at or datetime.now(timezone.utc)
    frames = parse_cao_excel(content, column_map, sheet=sheet)
    return [(df, _source_for(sid, when)) for sid, df in frames.items()]


def fetch_cao_cci(
    content: bytes,
    *,
    sheet: str | None = None,
    retrieved_at: datetime | None = None,
) -> list[tuple[pd.DataFrame, Source]]:
    """消費動向調査 Excel から 消費者態度指数＋構成DI を返す。"""
    return _fetch(content, CCI_COLUMN_MAP, sheet=sheet, retrieved_at=retrieved_at)


def fetch_cao_watcher(
    content: bytes,
    *,
    sheet: str | None = None,
    retrieved_at: datetime | None = None,
) -> list[tuple[pd.DataFrame, Source]]:
    """景気ウォッチャー調査 Excel から 現状・先行きDI を返す。"""
    return _fetch(content, WATCHER_COLUMN_MAP, sheet=sheet, retrieved_at=retrieved_at)
